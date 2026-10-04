from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

# Create presentation
prs = Presentation()

# Standard layout indices:
# 0 = Title Slide
# 1 = Title and Content

# Slide 1: Title Slide
slide_layout = prs.slide_layouts[0]
slide = prs.slides.add_slide(slide_layout)
title = slide.shapes.title
subtitle = slide.placeholders[1]

title.text = "Implementasi Google Analytics pada Sistem Jasa Website"
subtitle.text = "Optimalisasi Pelacakan Data Pengguna dan Transaksi E-Commerce"

# Helper function to add bullets
def add_bullets(tf, points):
    for i, point in enumerate(points):
        p = tf.add_paragraph() if i > 0 else tf.paragraphs[0]
        p.text = point
        p.level = 0
        p.font.size = Pt(20)

# Slide 2: Pendahuluan
slide_layout = prs.slide_layouts[1]
slide = prs.slides.add_slide(slide_layout)
title = slide.shapes.title
title.text = "Google Analytics 4 (GA4) & Tujuan"
tf = slide.placeholders[1].text_frame
bullets = [
    "Layanan analitik tingkat lanjut untuk menganalisis aktivitas dan perilaku pengunjung secara mendalam.",
    "Memantau jumlah dan demografi pengunjung website Jasa Website.",
    "Memetakan perjalanan pengguna (User Journey) dari melihat paket hingga selesai transaksi.",
    "Mengevaluasi performa sistem checkout dan halaman instruksi pembayaran.",
    "Membantu pengambilan keputusan bisnis berdasarkan data riil (Data-Driven)."
]
add_bullets(tf, bullets)

# Slide 3: Implementasi & Teknologi
slide_layout = prs.slide_layouts[1]
slide = prs.slides.add_slide(slide_layout)
title = slide.shapes.title
title.text = "Implementasi & Teknologi yang Digunakan"
tf = slide.placeholders[1].text_frame
bullets = [
    "Integrasi Dasar: Pembuatan akun, properti GA4, dan pemasangan tracking code secara global.",
    "Fitur Lanjutan (E-Commerce Tracking): Melacak seluruh corong konversi (Funnel), bukan sekadar lalu lintas biasa.",
    "Event Tracking Spesifik: Melacak interaksi di setiap tombol dan halaman kritis.",
    "Session Flash Tracking: Teknologi integrasi Backend (Laravel) ke Frontend untuk merekam transaksi valid.",
    "Sistem mentransmisikan data real-time ke server Google tanpa mengganggu kecepatan loading website."
]
add_bullets(tf, bullets)

# Slide 4: Pemetaan Event & Kriteria Analisis
slide_layout = prs.slide_layouts[1]
slide = prs.slides.add_slide(slide_layout)
title = slide.shapes.title
title.text = "Pemetaan Event Listener (The Funnel)"
tf = slide.placeholders[1].text_frame
bullets = [
    "view_item_list & select_item: Terdeteksi saat pengunjung melihat daftar paket dan klik tombol Order.",
    "begin_checkout: Otomatis terekam saat pengguna masuk ke halaman detail Checkout.",
    "add_payment_info: Dilacak saat pengguna membaca instruksi transfer pembayaran manual.",
    "purchase (Pembelian): Tereksekusi hanya saat pengguna sukses mengunggah bukti transfer (Data Revenue 100% valid)."
]
add_bullets(tf, bullets)

# Slide 5: Manfaat Analisis
slide_layout = prs.slide_layouts[1]
slide = prs.slides.add_slide(slide_layout)
title = slide.shapes.title
title.text = "Manfaat Analisis dari Data E-Commerce"
tf = slide.placeholders[1].text_frame
bullets = [
    "Cart Abandonment Rate: Mengidentifikasi persentase pengunjung yang memulai checkout namun batal membayar.",
    "Performa Paket: Membandingkan paket mana yang paling sering dilihat dengan yang paling banyak dibeli.",
    "Optimalisasi UX: Mengevaluasi efektivitas halaman instruksi pembayaran.",
    "Deteksi Kebocoran: Menemukan halaman mana yang menyebabkan pengunjung pergi (Bounce Rate tinggi)."
]
add_bullets(tf, bullets)

# Slide 6: Kesimpulan & Tindak Lanjut
slide_layout = prs.slide_layouts[1]
slide = prs.slides.add_slide(slide_layout)
title = slide.shapes.title
title.text = "Kesimpulan & Langkah ke Depan"
tf = slide.placeholders[1].text_frame
bullets = [
    "Sistem Jasa Website kini sepenuhnya cerdas dalam melacak aktivitas pengguna dan pendapatan transaksi.",
    "Solusi Session Flash telah berhasil menuntaskan celah tracking di sistem transfer manual.",
    "Tindak Lanjut 1: Pantau dasbor Analytics secara rutin (ingat ada proses sinkronisasi data 24-48 jam).",
    "Tindak Lanjut 2: Evaluasi sumber traffic (sosial media, iklan, SEO) mana yang menyumbang purchase terbanyak.",
    "Tindak Lanjut 3: Gunakan insight audiens untuk menciptakan strategi pemasaran yang lebih tertarget."
]
add_bullets(tf, bullets)

# Save presentation
prs.save('Presentasi_GA4_JasaWebsite.pptx')
print("Presentation generated successfully!")
