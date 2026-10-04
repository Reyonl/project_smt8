from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
import os

# Create presentation
prs = Presentation()

# Constants
TITLE_COLOR = RGBColor(31, 73, 125) # Professional Blue
TEXT_COLOR = RGBColor(64, 64, 64)   # Dark Gray
FONT_NAME = 'Calibri'

# Image paths
img1_path = r'C:\Users\ThinkPad\.gemini\antigravity\brain\b72d4c84-3b0f-450e-b63e-038b8a4622e0\tech_analytics_illus_1780044086337.png'
img2_path = r'C:\Users\ThinkPad\.gemini\antigravity\brain\b72d4c84-3b0f-450e-b63e-038b8a4622e0\visitor_chart_1780044102076.png'
img3_path = r'C:\Users\ThinkPad\.gemini\antigravity\brain\b72d4c84-3b0f-450e-b63e-038b8a4622e0\ga_dashboard_mockup_1780044118245.png'

def style_text_frame(tf):
    tf.word_wrap = True
    for paragraph in tf.paragraphs:
        paragraph.font.name = FONT_NAME
        paragraph.font.color.rgb = TEXT_COLOR
        paragraph.font.size = Pt(16)
        if paragraph.level == 0:
            paragraph.font.size = Pt(18)
        
def style_title(shape):
    shape.text_frame.paragraphs[0].font.name = FONT_NAME
    shape.text_frame.paragraphs[0].font.color.rgb = TITLE_COLOR
    shape.text_frame.paragraphs[0].font.bold = True

# -----------------
# SLIDE 1: PENDAHULUAN
# -----------------
slide = prs.slides.add_slide(prs.slide_layouts[1]) # Title and Content
title_shape = slide.shapes.title
title_shape.text = "1. Pendahuluan"
style_title(title_shape)

tf = slide.placeholders[1].text_frame
bullets = [
    "Pengertian Google Analytics:",
    "  • Layanan analitik web gratis dari Google untuk melacak dan melaporkan lalu lintas website.",
    "Tujuan Penggunaan pada Website Antigravity:",
    "  • Mengukur efektivitas website sebagai sarana penawaran jasa pembuatan website.",
    "  • Mengidentifikasi target pasar potensial secara lebih presisi.",
    "Pentingnya Analisis Data Pengunjung:",
    "  • Memahami perilaku dan minat audiens secara nyata (Data-Driven).",
    "  • Mengurangi risiko pengambilan keputusan yang salah dalam operasional bisnis."
]
for i, point in enumerate(bullets):
    p = tf.add_paragraph() if i > 0 else tf.paragraphs[0]
    p.text = point
    p.level = 0 if not point.startswith("  •") else 1
style_text_frame(tf)

# Adjust text box width to make room for image
slide.placeholders[1].width = Inches(5.5)

# Insert Image 1 (Illustration) on the right
if os.path.exists(img1_path):
    slide.shapes.add_picture(img1_path, Inches(6.0), Inches(2.0), width=Inches(3.5))

# -----------------
# SLIDE 2: IMPLEMENTASI
# -----------------
slide = prs.slides.add_slide(prs.slide_layouts[1])
title_shape = slide.shapes.title
title_shape.text = "2. Implementasi Google Analytics"
style_title(title_shape)

tf = slide.placeholders[1].text_frame
bullets = [
    "Cara Integrasi ke Website Antigravity:",
    "  • Pembuatan akun dan pengaturan Properti (Web Stream) di GA4.",
    "Tahapan Pemasangan Tracking Code:",
    "  • Menambahkan script gtag.js ke dalam tag <head> di seluruh halaman (global).",
    "  • Melakukan verifikasi melalui Realtime Report atau Google Tag Assistant.",
    "Fitur Analytics yang Digunakan:",
    "  • Traffic Acquisition, E-Commerce Events, dan Event Tracking Custom.",
    "Cara Pengumpulan Data:",
    "  • Sistem menggunakan Event Listeners (misal: 'purchase' / 'view_item_list') yang ditembakkan melalui browser pengunjung."
]
for i, point in enumerate(bullets):
    p = tf.add_paragraph() if i > 0 else tf.paragraphs[0]
    p.text = point
    p.level = 0 if not point.startswith("  •") else 1
style_text_frame(tf)


# -----------------
# SLIDE 3: HASIL DAN KRITERIA ANALISIS
# -----------------
slide = prs.slides.add_slide(prs.slide_layouts[1])
title_shape = slide.shapes.title
title_shape.text = "3. Hasil dan Kriteria Analisis"
style_title(title_shape)

# Create two columns for this slide by adjusting width
tf = slide.placeholders[1].text_frame
bullets = [
    "Data Utama yang Dianalisis:",
    "  • Jumlah Pengunjung & Pengguna Aktif",
    "  • Lama Kunjungan & Bounce Rate",
    "  • Halaman Terpopuler (Pageviews)",
    "  • Sumber Traffic (Sosial media, Organik, dsb)",
    "  • Demografi: Lokasi & Device (Mobile/Desktop)",
    "  • Klik & Interaksi Spesifik (Tombol Pesan)",
    "Manfaat bagi Antigravity:",
    "  • Optimalisasi antarmuka berdasarkan device pengguna.",
    "  • Penyesuaian materi promosi dengan halaman terpopuler."
]
for i, point in enumerate(bullets):
    p = tf.add_paragraph() if i > 0 else tf.paragraphs[0]
    p.text = point
    p.level = 0 if not point.startswith("  •") else 1
style_text_frame(tf)

slide.placeholders[1].width = Inches(5.0)

# Insert Image 2 (Chart) on the right
if os.path.exists(img2_path):
    slide.shapes.add_picture(img2_path, Inches(5.2), Inches(2.0), width=Inches(4.5))

# -----------------
# SLIDE 4: KESIMPULAN
# -----------------
slide = prs.slides.add_slide(prs.slide_layouts[1])
title_shape = slide.shapes.title
title_shape.text = "4. Kesimpulan & Rekomendasi"
style_title(title_shape)

tf = slide.placeholders[1].text_frame
bullets = [
    "Kesimpulan Implementasi:",
    "  • Google Analytics telah terintegrasi sukses dan berhasil memetakan seluruh interaksi pengguna di platform.",
    "Dampak Analisis terhadap Sistem:",
    "  • Memberikan kejelasan titik kelemahan sistem (misal: tingginya pembatalan saat checkout) untuk perbaikan instan.",
    "Saran Pengembangan:",
    "  • Menjalankan evaluasi rutin bulanan berdasarkan laporan analitik.",
    "  • Menggunakan A/B Testing untuk halaman dengan Bounce Rate tinggi.",
    "  • Fokus kampanye iklan pada sumber traffic yang paling banyak melakukan konversi (purchase)."
]
for i, point in enumerate(bullets):
    p = tf.add_paragraph() if i > 0 else tf.paragraphs[0]
    p.text = point
    p.level = 0 if not point.startswith("  •") else 1
style_text_frame(tf)

slide.placeholders[1].width = Inches(5.0)

# Insert Image 3 (Dashboard Mockup) on the right
if os.path.exists(img3_path):
    slide.shapes.add_picture(img3_path, Inches(5.5), Inches(2.0), width=Inches(4.0))

# Save presentation
output_path = r'c:\laragon\www\jasa-website\Presentasi_Mahasiswa_Antigravity.pptx'
prs.save(output_path)
print(f"Presentation saved to {output_path}")
