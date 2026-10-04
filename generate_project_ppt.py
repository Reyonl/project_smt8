from pptx import Presentation
from pptx.util import Inches, Pt

# Create presentation
prs = Presentation()

# Slide 1: Title Slide
slide = prs.slides.add_slide(prs.slide_layouts[0])
title = slide.shapes.title
subtitle = slide.placeholders[1]
title.text = "Jasa Website"
subtitle.text = "Platform Pemesanan & Manajemen Proyek Website Terpadu\nRingkasan Keseluruhan Proyek"

def add_bullets(tf, points):
    for i, point in enumerate(points):
        p = tf.add_paragraph() if i > 0 else tf.paragraphs[0]
        p.text = point
        p.level = 0
        p.font.size = Pt(18)

# Slide 2: Latar Belakang & Tujuan
slide = prs.slides.add_slide(prs.slide_layouts[1])
slide.shapes.title.text = "Latar Belakang & Tujuan"
bullets = [
    "Digitalisasi Pemesanan: Mempermudah klien dalam memilih dan memesan layanan pembuatan website.",
    "Transparansi Proyek: Klien dapat memantau status pesanan dan progress pengerjaan secara real-time.",
    "Sentralisasi Komunikasi: Menggantikan komunikasi terpisah (chat) dengan Ruang Diskusi yang terintegrasi di dalam sistem.",
    "Keputusan Berbasis Data: Melacak performa bisnis secara akurat menggunakan sistem Analytics lanjutan."
]
add_bullets(slide.placeholders[1].text_frame, bullets)

# Slide 3: Fitur Unggulan (Sisi Klien/User)
slide = prs.slides.add_slide(prs.slide_layouts[1])
slide.shapes.title.text = "Fitur Unggulan (Sisi Klien)"
bullets = [
    "Katalog Paket Dinamis: Klien dapat melihat pilihan paket website beserta harganya dengan tampilan modern.",
    "Custom Order: Fasilitas form khusus bagi klien yang menginginkan website dengan spesifikasi unik.",
    "Sistem Pembayaran Manual: Alur upload bukti transfer yang mudah dan terstruktur.",
    "Ruang Diskusi & Progress: Klien dapat berinteraksi langsung dengan admin dan melihat riwayat pengerjaan proyek (Timeline)."
]
add_bullets(slide.placeholders[1].text_frame, bullets)

# Slide 4: Fitur Unggulan (Sisi Admin)
slide = prs.slides.add_slide(prs.slide_layouts[1])
slide.shapes.title.text = "Fitur Unggulan (Sisi Admin)"
bullets = [
    "Manajemen Paket: Admin memiliki kendali penuh untuk menambah, mengubah, atau menghapus paket layanan.",
    "Verifikasi Pembayaran: Sistem validasi bukti transfer dengan sekali klik (Approve / Reject).",
    "Update Progress: Kemampuan admin untuk memperbarui status pesanan dan memberikan catatan progress harian ke klien.",
    "Laporan (Reporting): Fitur rekapitulasi data pesanan yang dapat dicetak (Printable) untuk evaluasi bulanan."
]
add_bullets(slide.placeholders[1].text_frame, bullets)

# Slide 5: Alur Pemesanan (User Journey)
slide = prs.slides.add_slide(prs.slide_layouts[1])
slide.shapes.title.text = "Alur Kerja Sistem (User Journey)"
bullets = [
    "1. Pemilihan Produk: Klien memilih paket reguler atau mengajukan penawaran custom.",
    "2. Checkout & Bayar: Klien menerima kode pesanan (Invoice) dan mengunggah bukti transfer bank.",
    "3. Verifikasi & Pengerjaan: Admin menyetujui pembayaran dan sistem mengubah status menjadi 'Diproses'.",
    "4. Interaksi Proyek: Klien dan Admin saling berkomunikasi dan berbagi lampiran melalui Ruang Diskusi.",
    "5. Penyerahan Final: Pesanan ditandai selesai (Completed) dan data dikunci."
]
add_bullets(slide.placeholders[1].text_frame, bullets)

# Slide 6: Integrasi Canggih (Analytics)
slide = prs.slides.add_slide(prs.slide_layouts[1])
slide.shapes.title.text = "Integrasi Canggih & Analytics"
bullets = [
    "Google Analytics 4 (GA4) terpasang di seluruh alur penting aplikasi.",
    "E-Commerce Funnel Tracking: Sistem merekam data mulai dari klien melihat paket (view_item) hingga sukses membayar (purchase).",
    "Session Flash Tracking: Solusi inovatif yang menjamin data pendapatan (Revenue) 100% valid terekam ke Google saat bukti pembayaran diunggah.",
    "Dampak: Membantu tim manajemen mengevaluasi paket mana yang paling laris dan dari mana asal konversi terbesar."
]
add_bullets(slide.placeholders[1].text_frame, bullets)

# Slide 7: Kesimpulan
slide = prs.slides.add_slide(prs.slide_layouts[1])
slide.shapes.title.text = "Kesimpulan Keseluruhan"
bullets = [
    "Aplikasi Jasa Website adalah solusi 'End-to-End' yang sangat komprehensif.",
    "Memadukan User Experience (UX) yang menarik di sisi klien dengan User Interface (UI) fungsional di sisi admin.",
    "Dibangun dengan teknologi modern (Laravel) yang aman, cepat, dan terukur.",
    "Siap mendukung pertumbuhan bisnis melalui pencatatan data yang presisi dan pelacakan analitik pintar."
]
add_bullets(slide.placeholders[1].text_frame, bullets)

# Save presentation
prs.save('Presentasi_Keseluruhan_JasaWebsite.pptx')
print("Overall Project Presentation generated successfully!")
